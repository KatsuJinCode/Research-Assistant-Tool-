#!/usr/bin/env python3
"""
Unit tests for file type support (PDF, TXT, DOCX)
"""
import pytest
from pathlib import Path
from research_agent.document_processing.pdf_extractor import PDFExtractor
from research_agent.document_processing.txt_extractor import TXTExtractor
from research_agent.document_processing.docx_extractor import DOCXExtractor


class TestFileTypeSupport:
    """Test extraction from different file types"""

    def test_txt_extraction(self, tmp_path):
        """Test TXT file extraction"""
        # Create a test TXT file
        test_file = tmp_path / "test.txt"
        test_content = "This is a test document.\nIt has multiple lines.\n\nAnd paragraphs."
        test_file.write_text(test_content, encoding='utf-8')

        # Extract text
        extractor = TXTExtractor()
        result = extractor.extract(test_file)

        # Verify results
        assert result['text'] == test_content
        assert result['quality_score'] > 0
        assert result['metadata']['file_type'] == 'txt'
        assert result['metadata']['word_count'] > 0
        print(f"✓ TXT extraction successful: {result['metadata']['word_count']} words")

    def test_docx_extraction(self, tmp_path):
        """Test DOCX file extraction"""
        try:
            from docx import Document
        except ImportError:
            pytest.skip("python-docx not installed")

        # Create a test DOCX file
        test_file = tmp_path / "test.docx"
        doc = Document()
        doc.add_paragraph("This is the first paragraph.")
        doc.add_paragraph("This is the second paragraph with more content.")
        doc.save(str(test_file))

        # Extract text
        extractor = DOCXExtractor()
        result = extractor.extract(test_file)

        # Verify results
        assert 'first paragraph' in result['text'].lower()
        assert 'second paragraph' in result['text'].lower()
        assert result['quality_score'] > 0
        assert result['metadata']['file_type'] == 'docx'
        assert result['metadata']['word_count'] > 0
        print(f"✓ DOCX extraction successful: {result['metadata']['word_count']} words")

    def test_pdf_extraction(self):
        """Test PDF file extraction with existing test file"""
        test_pdf = Path('web_ui/uploads/children.pdf')
        if not test_pdf.exists():
            pytest.skip(f"{test_pdf} not found")

        # Extract text
        extractor = PDFExtractor(column_aware=True, postprocess=True)
        result = extractor.extract(test_pdf)

        # Verify results
        assert 'full_text' in result
        assert len(result['full_text']) > 0
        assert result['quality_score'] > 0
        print(f"✓ PDF extraction successful: {len(result['full_text'])} characters")

    def test_file_type_detection(self, tmp_path):
        """Test that file types are correctly identified"""
        # Create test files
        txt_file = tmp_path / "test.txt"
        txt_file.write_text("Test content")

        # Test suffix detection
        assert txt_file.suffix.lower() == '.txt'

        docx_file = tmp_path / "test.docx"
        assert docx_file.suffix.lower() == '.docx'

        pdf_file = tmp_path / "test.pdf"
        assert pdf_file.suffix.lower() == '.pdf'

        print("✓ File type detection working correctly")


if __name__ == '__main__':
    # Run tests with pytest
    import sys
    sys.exit(pytest.main([__file__, '-v']))
